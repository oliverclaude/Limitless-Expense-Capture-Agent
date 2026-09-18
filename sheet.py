"""Staging workbook: YYYY-MM-Expenses.xlsx in the month folder."""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

HEADERS = [
    "claim date",
    "personal",
    "company",
    "vat",
    "comment",
    "journal",
    "proof file",
]


def merchant_slug(merchant: str | None) -> str:
    if not merchant:
        return ""
    slug = re.sub(r"[^A-Za-z0-9]+", "-", merchant.strip())
    slug = re.sub(r"-{2,}", "-", slug).strip("-")
    return slug[:60]


def confirm_proofs(
    orig: Path,
    scan: Path | None,
    claim_date: str | None,
    merchant: str | None,
) -> tuple[Path, Path | None, str]:
    """Rename pending proofs to YYYY-MM-DD_<merchant>_orig/scan. Return proof filename."""
    slug = merchant_slug(merchant)
    if not slug:
        proof = scan or orig
        return orig, scan, proof.name
    day = claim_date if claim_date and re.fullmatch(r"\d{4}-\d{2}-\d{2}", claim_date) else orig.name[:10]
    new_orig = _rename_proof(orig, day, slug)
    new_scan = _rename_proof(scan, day, slug) if scan is not None else None
    proof = new_scan or new_orig
    return new_orig, new_scan, proof.name


def month_sheet_path(month_dir: Path) -> Path:
    return month_dir / f"{month_dir.name}-Expenses.xlsx"


AMOUNT_TOLERANCE = 0.02


def find_duplicate(sheet_path: Path, fields: dict[str, Any]) -> dict[str, Any] | None:
    """Same claim date, merchant slug, and personal amount within R0.02."""
    if not sheet_path.is_file():
        return None
    date = str(fields.get("claim_date") or "").strip()
    slug = merchant_slug(fields.get("merchant"))
    amount = fields.get("personal_amount")
    proof = str(fields.get("proof_file") or "")
    if not date or not slug or amount is None:
        return None
    wb = load_workbook(sheet_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            return None
        for row in ws.iter_rows(min_row=2, values_only=True):
            existing_date = _as_day(row[0] if len(row) else None)
            existing_amount = row[1] if len(row) > 1 else None
            existing_proof = str(row[6] or "") if len(row) > 6 else ""
            if existing_proof and proof and existing_proof == proof:
                continue
            if existing_date != date:
                continue
            try:
                existing_num = float(existing_amount)
            except (TypeError, ValueError):
                continue
            if abs(existing_num - float(amount)) > AMOUNT_TOLERANCE:
                continue
            proof_slug = _slug_from_proof(existing_proof)
            if proof_slug.lower() == slug.lower():
                return {
                    "claim_date": existing_date,
                    "personal": existing_num,
                    "proof_file": existing_proof,
                    "merchant_slug": proof_slug,
                }
    finally:
        wb.close()
    return None


def _as_day(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text[:10]


def _slug_from_proof(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"_(orig|scan)(-\d+)?$", "", stem)
    parts = stem.split("_", 1)
    if len(parts) == 2:
        return parts[1]
    return stem


def write_claim_row(sheet_path: Path, fields: dict[str, Any]) -> Path:
    sheet_path.parent.mkdir(parents=True, exist_ok=True)
    if sheet_path.is_file():
        wb = load_workbook(sheet_path)
        ws = wb.active
        _ensure_header(ws)
    else:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(HEADERS)
    row = [
        fields.get("claim_date") or "",
        fields.get("personal_amount"),
        fields.get("company_amount"),
        fields.get("vat"),
        fields.get("comment") or "",
        fields.get("journal") or "",
        fields.get("proof_file") or "",
    ]
    proof = str(fields.get("proof_file") or "")
    _compact_used_rows(ws)
    existing = _row_for_proof(ws, proof)
    if existing:
        for col, value in enumerate(row, start=1):
            ws.cell(existing, col, value)
    else:
        dest = _last_used_row(ws) + 1
        for col, value in enumerate(row, start=1):
            ws.cell(dest, col, value)
    _atomic_save(wb, sheet_path)
    return sheet_path


def _proof_kind(path: Path) -> tuple[str, str]:
    stem = path.stem
    if "_scan-" in stem or stem.endswith("_scan"):
        return "scan", ".jpg" if path.suffix.lower() in {".jpg", ".jpeg"} else path.suffix
    return "orig", path.suffix


def _index_suffix(path: Path) -> str:
    match = re.search(r"_(?:orig|scan)-(\d+)$", path.stem)
    if match:
        return f"-{match.group(1)}"
    return ""


def _rename_proof(path: Path, day: str, slug: str) -> Path:
    kind, ext = _proof_kind(path)
    dest = path.with_name(f"{day}_{slug}_{kind}{_index_suffix(path)}{ext}")
    if dest == path:
        return path
    if dest.exists():
        dest.unlink()
    path.replace(dest)
    return dest


def _ensure_header(ws: Worksheet) -> None:
    if ws.max_row == 0 or all(c.value is None for c in ws[1]):
        for col, name in enumerate(HEADERS, start=1):
            ws.cell(1, col, name)


def _last_used_row(ws: Worksheet) -> int:
    for row in range(ws.max_row or 1, 0, -1):
        if any(ws.cell(row, col).value not in (None, "") for col in range(1, len(HEADERS) + 1)):
            return row
    return 1


def _trim_empty_tail(ws: Worksheet) -> None:
    last = _last_used_row(ws)
    if (ws.max_row or 1) > last:
        ws.delete_rows(last + 1, (ws.max_row or last) - last)


def _compact_used_rows(ws: Worksheet) -> None:
    used: list[list[Any]] = []
    for row in range(1, (ws.max_row or 1) + 1):
        vals = [ws.cell(row, col).value for col in range(1, len(HEADERS) + 1)]
        if any(v not in (None, "") for v in vals):
            used.append(vals)
    if not used:
        return
    if ws.max_row and ws.max_row > 1:
        ws.delete_rows(1, ws.max_row)
    for r, vals in enumerate(used, start=1):
        for c, value in enumerate(vals, start=1):
            ws.cell(r, c, value)


def _row_for_proof(ws: Worksheet, proof: str) -> int | None:
    if not proof:
        return None
    last = _last_used_row(ws)
    for row in range(2, last + 1):
        if str(ws.cell(row, 7).value or "") == proof:
            return row
    return None


def _atomic_save(wb: WorkbookType, dest: Path) -> None:
    fd, tmp = tempfile.mkstemp(prefix=".tmp-", suffix=".xlsx", dir=dest.parent)
    os.close(fd)
    tmp_path = Path(tmp)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, dest)
    except Exception:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise
